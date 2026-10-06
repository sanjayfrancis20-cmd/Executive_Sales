import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

os.makedirs('deliverables', exist_ok=True)

prs = Presentation()
# Set widescreen 16:9
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette Definitions
C_NAVY_DARK = RGBColor(15, 23, 42)      # #0F172A
C_NAVY_BLUE = RGBColor(30, 58, 138)     # #1E3A8A
C_ACCENT_BLUE = RGBColor(37, 99, 235)   # #2563EB
C_EMERALD = RGBColor(16, 185, 129)      # #10B981
C_CRIMSON = RGBColor(239, 68, 68)       # #EF4444
C_SLATE_GRAY = RGBColor(100, 116, 139)  # #64748B
C_LIGHT_BG = RGBColor(248, 250, 252)    # #F8FAFC
C_CARD_BG = RGBColor(255, 255, 255)     # #FFFFFF
C_TEXT_DARK = RGBColor(30, 41, 59)      # #1E293B
C_TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B
C_WHITE = RGBColor(255, 255, 255)

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_text="EXECUTIVE PERFORMANCE DASHBOARD"):
    # Header container
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p_cat = tf.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(9.5)
    p_cat.font.bold = True
    p_cat.font.color.rgb = C_ACCENT_BLUE
    
    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.size = Pt(20)
    p_title.font.bold = True
    p_title.font.color.rgb = C_NAVY_DARK

def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.2)
    else:
        shape.line.fill.background()
    return shape

# ==========================================
# SLIDE 1: TITLE SLIDE (Dark Premium Theme)
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = C_NAVY_DARK
bg1.line.fill.background()

# Title text box
tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(3.8))
tf1 = tb1.text_frame
tf1.word_wrap = True

p_badge = tf1.paragraphs[0]
p_badge.text = "DATA ANALYST INTERNSHIP • TASK 4 CAPSTONE"
p_badge.font.size = Pt(12)
p_badge.font.bold = True
p_badge.font.color.rgb = RGBColor(96, 165, 250)

p_main = tf1.add_paragraph()
p_main.text = "Executive Sales & Financial Performance Dashboard"
p_main.font.size = Pt(32)
p_main.font.bold = True
p_main.font.color.rgb = C_WHITE
p_main.space_before = Pt(12)

p_sub = tf1.add_paragraph()
p_sub.text = "Interactive Stakeholder Analytics, Trend Forecasting, Profit Margin Optimization, and Tableau BI Architecture"
p_sub.font.size = Pt(15)
p_sub.font.color.rgb = RGBColor(203, 213, 225)
p_sub.space_before = Pt(14)

# Metadata card on bottom
add_card(s1, Inches(1.2), Inches(5.4), Inches(10.9), Inches(1.2), bg_color=RGBColor(30, 41, 59))
meta_tb = s1.shapes.add_textbox(Inches(1.5), Inches(5.5), Inches(10.3), Inches(1.0))
mtf = meta_tb.text_frame
mtf.word_wrap = True
p_meta = mtf.paragraphs[0]
p_meta.text = "Author: Data Analyst Intern   |   Platform: Tableau Public / Power BI   |   Dataset: Kaggle Sales & Financial Records (9,994 Rows)\nDeliverables: Interactive Dashboard, Business Summary Deck, Interview Solutions"
p_meta.font.size = Pt(11)
p_meta.font.color.rgb = RGBColor(226, 232, 240)

# ==========================================
# SLIDE 2: BUSINESS CONTEXT & DASHBOARD OBJECTIVES
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Business Problem & Stakeholder Objectives", "PROJECT FOUNDATION")

cards_data = [
    ("Executive Problem Statement", 
     "Commercial stakeholders lacked unified visibility across regional sales, product margins, and discount leakages. Top-line revenue growth was masking severe margin erosion in specific sub-categories and territories.",
     C_ACCENT_BLUE),
    ("Core Dashboard Purpose", 
     "Deliver an intuitive, high-impact Tableau dashboard that enables C-suite executives, VP of Sales, and Regional Directors to evaluate KPIs, explore trends, drill into problem areas, and take strategic action.",
     C_EMERALD),
    ("Key Stakeholder Requirements", 
     "• High-level KPI Scorecards for rapid executive scanning\n• Interactive Slicers for granular multi-dimensional filtering\n• Monthly & Annual Time-Series trend analysis\n• Color-coded profitability matrix to isolate loss-making units",
     C_NAVY_BLUE)
]

for idx, (title, desc, accent) in enumerate(cards_data):
    left = Inches(0.8 + idx * 3.95)
    add_card(s2, left, Inches(1.6), Inches(3.8), Inches(5.1), bg_color=RGBColor(241, 245, 249), border_color=accent)
    tb = s2.shapes.add_textbox(left + Inches(0.25), Inches(1.85), Inches(3.3), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = accent
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = C_TEXT_DARK
    p2.space_before = Pt(12)

# ==========================================
# SLIDE 3: EXECUTIVE KPI SCORECARD
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "Executive KPI Scorecard & Summary Cards", "HIGH-LEVEL PERFORMANCE")

kpi_cards = [
    ("TOTAL REVENUE", "$13.05M", "+3.2% vs Target", C_ACCENT_BLUE),
    ("TOTAL NET PROFIT", "$1.39M", "10.69% Net Margin", C_EMERALD),
    ("TOTAL ORDERS", "4,998", "30,282 Units Sold", C_NAVY_DARK),
    ("AVG ORDER VALUE", "$2,611", "Healthy B2B Basket", C_SLATE_GRAY)
]

for idx, (lbl, val, sub, col) in enumerate(kpi_cards):
    left = Inches(0.8 + idx * 2.95)
    add_card(s3, left, Inches(1.5), Inches(2.85), Inches(1.8), bg_color=RGBColor(248, 250, 252), border_color=col)
    tb = s3.shapes.add_textbox(left + Inches(0.2), Inches(1.65), Inches(2.45), Inches(1.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = lbl
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = C_TEXT_MUTED
    
    p1 = tf.add_paragraph()
    p1.text = val
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = col
    
    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = C_TEXT_MUTED

# Summary table / observations
add_card(s3, Inches(0.8), Inches(3.6), Inches(11.7), Inches(3.2), bg_color=RGBColor(255, 255, 255), border_color=RGBColor(226, 232, 240))
tb_table = s3.shapes.add_textbox(Inches(1.1), Inches(3.8), Inches(11.1), Inches(2.8))
tft = tb_table.text_frame
tft.word_wrap = True

pt_title = tft.paragraphs[0]
pt_title.text = "Key Observations Across Multi-Year Operations (2021 - 2024)"
pt_title.font.size = Pt(14)
pt_title.font.bold = True
pt_title.font.color.rgb = C_NAVY_DARK

bullets = [
    "Revenue Resilience: Annual sales have maintained steady volume around ~$3.2M - $3.4M per annum.",
    "Margin Pressure in 2024: Net profit margin dipped from 11.34% in 2023 to 9.67% in 2024 due to aggressive discounting across select categories.",
    "Volume vs Profit Realization: Average Order Value remains robust at $2,611 across 4,998 commercial orders, indicating strong transaction sizing.",
    "Optimization Potential: Substantial bottom-line improvement can be captured by plugging negative-margin product categories without reducing sales volume."
]

for b in bullets:
    p = tft.add_paragraph()
    p.text = "• " + b
    p.font.size = Pt(11)
    p.font.color.rgb = C_TEXT_DARK
    p.space_before = Pt(8)

# ==========================================
# SLIDE 4: TIME-SERIES TREND ANALYSIS
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "Time-Series Trend & Seasonality Analysis", "TEMPORAL PATTERNS")

if os.path.exists('assets/chart_timeseries_trend.png'):
    s4.shapes.add_picture('assets/chart_timeseries_trend.png', Inches(0.8), Inches(1.5), Inches(7.5), Inches(5.3))

add_card(s4, Inches(8.5), Inches(1.5), Inches(4.0), Inches(5.3), bg_color=RGBColor(241, 245, 249))
tb_ts = s4.shapes.add_textbox(Inches(8.7), Inches(1.7), Inches(3.6), Inches(4.9))
tfts = tb_ts.text_frame
tfts.word_wrap = True

p_h = tfts.paragraphs[0]
p_h.text = "Key Time-Series Insights"
p_h.font.size = Pt(14)
p_h.font.bold = True
p_h.font.color.rgb = C_NAVY_DARK

ts_notes = [
    "Q4 Seasonal Spike: Monthly revenue consistently peaks in November & December driven by corporate budget spend and holiday inventory restocking.",
    "Profit Alignment: Profit closely tracks top-line sales during Q1-Q3 but experiences margin compression in promotional periods due to discount escalation.",
    "Cyclical Mid-Year Dips: February and July exhibit recurring dips across all four operating years, suggesting ideal windows for marketing campaigns and inventory planning.",
    "Forecasting Implication: Strong seasonality enables 90-day rolling moving averages to project forward demand with high reliability."
]
for n in ts_notes:
    p = tfts.add_paragraph()
    p.text = "• " + n
    p.font.size = Pt(10)
    p.font.color.rgb = C_TEXT_DARK
    p.space_before = Pt(10)

# ==========================================
# SLIDE 5: REGIONAL PERFORMANCE BREAKDOWN
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Geographic Performance: Sales & Margin Variance", "REGIONAL INSIGHTS")

if os.path.exists('assets/chart_regional_performance.png'):
    s5.shapes.add_picture('assets/chart_regional_performance.png', Inches(0.8), Inches(1.5), Inches(7.5), Inches(5.3))

add_card(s5, Inches(8.5), Inches(1.5), Inches(4.0), Inches(5.3), bg_color=RGBColor(241, 245, 249))
tb_reg = s5.shapes.add_textbox(Inches(8.7), Inches(1.7), Inches(3.6), Inches(4.9))
tfreg = tb_reg.text_frame
tfreg.word_wrap = True

p_rh = tfreg.paragraphs[0]
p_rh.text = "Regional Diagnostic"
p_rh.font.size = Pt(14)
p_rh.font.bold = True
p_rh.font.color.rgb = C_NAVY_DARK

reg_notes = [
    "West Region Dominance: Generates the highest total sales volume ($4.1M+) and maintains the healthiest net profit margin (>11.5%).",
    "East Region Strength: Solid contributor with strong enterprise tech sales across New York and Massachusetts hubs.",
    "Central Region Margin Leak: Despite generating over $3.0M in sales, Central suffers margin dilution due to aggressive price-matching and excessive discounting on furniture.",
    "South Region Opportunity: Lowest volume contribution (~16%), but healthy margins suggest untapped headroom for sales rep expansion."
]
for n in reg_notes:
    p = tfreg.add_paragraph()
    p.text = "• " + n
    p.font.size = Pt(10)
    p.font.color.rgb = C_TEXT_DARK
    p.space_before = Pt(10)

# ==========================================
# SLIDE 6: PRODUCT SUB-CATEGORY PROFITABILITY
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "Product Sub-Category Profitability Matrix", "PORTFOLIO ANALYSIS")

if os.path.exists('assets/chart_subcategory_profit.png'):
    s6.shapes.add_picture('assets/chart_subcategory_profit.png', Inches(0.8), Inches(1.5), Inches(7.5), Inches(5.3))

add_card(s6, Inches(8.5), Inches(1.5), Inches(4.0), Inches(5.3), bg_color=RGBColor(241, 245, 249))
tb_p = s6.shapes.add_textbox(Inches(8.7), Inches(1.7), Inches(3.6), Inches(4.9))
tfp = tb_p.text_frame
tfp.word_wrap = True

p_ph = tfp.paragraphs[0]
p_ph.text = "Profit Engine vs Leaks"
p_ph.font.size = Pt(14)
p_ph.font.bold = True
p_ph.font.color.rgb = C_NAVY_DARK

prod_notes = [
    "Profit Champions: Copiers, Phones, and Accessories represent over 55% of total company net profit with stellar margins (>20%).",
    "Office Supplies Stability: Paper, Binders, and Storage generate consistent positive cash flow with minimal price sensitivity.",
    "Severe Drag from Tables: Tables generate negative total profit across regions due to deep discounting (30-50%) combined with high fulfillment freight costs.",
    "Strategic Fix: Implement minimum floor pricing and eliminate discounts over 15% on freight-heavy furniture units."
]
for n in prod_notes:
    p = tfp.add_paragraph()
    p.text = "• " + n
    p.font.size = Pt(10)
    p.font.color.rgb = C_TEXT_DARK
    p.space_before = Pt(10)

# ==========================================
# SLIDE 7: CUSTOMER SEGMENT BREAKDOWN
# ==========================================
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "Customer Segment Share & Discount Sensitivity", "SEGMENT ANALYTICS")

if os.path.exists('assets/chart_segment_distribution.png'):
    s7.shapes.add_picture('assets/chart_segment_distribution.png', Inches(0.8), Inches(1.5), Inches(5.5), Inches(5.3))

add_card(s7, Inches(6.6), Inches(1.5), Inches(5.9), Inches(5.3), bg_color=RGBColor(241, 245, 249))
tb_seg = s7.shapes.add_textbox(Inches(6.9), Inches(1.7), Inches(5.3), Inches(4.9))
tfseg = tb_seg.text_frame
tfseg.word_wrap = True

p_sh = tfseg.paragraphs[0]
p_sh.text = "Segment Profitability Dynamics"
p_sh.font.size = Pt(14)
p_sh.font.bold = True
p_sh.font.color.rgb = C_NAVY_DARK

seg_insights = [
    "Consumer Segment (52.2% Sales): The largest revenue contributor. Higher transaction frequency but higher discount utilization.",
    "Corporate Segment (30.1% Sales): The most stable, high-margin segment. Corporate contracts exhibit lower price elasticity and repeat purchasing behavior.",
    "Home Office (17.7% Sales): The fastest growing post-pandemic segment. High demand for Technology and ergonomic Office Supplies with healthy margins.",
    "Key Takeaway: Prioritize Corporate enterprise multi-year contracts while curtailing consumer flash discounts."
]
for n in seg_insights:
    p = tfseg.add_paragraph()
    p.text = "• " + n
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_TEXT_DARK
    p.space_before = Pt(10)

# ==========================================
# SLIDE 8: TABLEAU DASHBOARD ARCHITECTURE
# ==========================================
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "Interactive Dashboard Architecture & Design System", "TECHNICAL IMPLEMENTATION")

arch_sections = [
    ("1. Core KPI Metrics & Cards", 
     "• Total Sales ($13.05M)\n• Net Profit ($1.39M)\n• Profit Margin (10.69%)\n• YoY Growth % and Average Order Value\n• Real-time conditional KPI color indicators"),
    ("2. Slicers & Interactivity", 
     "• Global Region Slicer (West, East, Central, South)\n• Category & Sub-Category multi-select filter\n• Year & Timeline drilldown (2021-2024)\n• Interactive cross-filtering via Worksheet Actions"),
    ("3. Advanced Tableau Logic", 
     "• Fixed LOD Expressions for Region Benchmark:\n  { FIXED [Region] : SUM([Sales]) }\n• Dynamic Profit Status: IF [Profit] > 0 THEN 'Profitable' ELSE 'Loss' END\n• Parameter Actions for Metric Switching"),
    ("4. Visual Hierarchy & Theme", 
     "• Cohesive Slate/Navy UI with 8px grid alignment\n• Semantic color palette (Emerald profit, Crimson loss)\n• Integrated top navigation tabs for instant view switching\n• High-readability typography and tooltips")
]

for idx, (title, content) in enumerate(arch_sections):
    row = idx // 2
    col = idx % 2
    left = Inches(0.8 + col * 5.95)
    top = Inches(1.6 + row * 2.7)
    
    add_card(s8, left, top, Inches(5.75), Inches(2.45), bg_color=RGBColor(248, 250, 252), border_color=C_ACCENT_BLUE)
    tb = s8.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(5.35), Inches(2.1))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p_t = tf.paragraphs[0]
    p_t.text = title
    p_t.font.size = Pt(12.5)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK
    
    p_c = tf.add_paragraph()
    p_c.text = content
    p_c.font.size = Pt(9.5)
    p_c.font.color.rgb = C_TEXT_DARK
    p_c.space_before = Pt(6)

# ==========================================
# SLIDE 9: STRATEGIC BUSINESS RECOMMENDATIONS
# ==========================================
s9 = prs.slides.add_slide(blank_layout)
add_header(s9, "Actionable Business Recommendations", "STRATEGIC ROADMAP")

recs = [
    ("Cap Maximum Discounts at 15%", 
     "Data indicates discount rates exceeding 20% invariably convert positive gross margins into negative net profits, especially in Tables and Bookcases. Instituting an automated approval gate will salvage an estimated $120K in annual profit.",
     C_CRIMSON),
    ("Double Down on High-Margin Technology", 
     "Copiers and Phones generate >25% profit margin with high average ticket sizes. Reallocate marketing budgets toward corporate technology promotions to elevate portfolio margins.",
     C_EMERALD),
    ("Central Region Turnaround Strategy", 
     "Perform a deep logistics and pricing audit on the Central territory to rectify freight cost overheads and align regional discounting policies with the West territory model.",
     C_ACCENT_BLUE),
    ("Expand Corporate Contract Retention", 
     "Corporate clients yield superior customer lifetime value (CLV) and lower discount sensitivity. Introduce tiered loyalty rebate structures rather than upfront transaction discounts.",
     C_NAVY_DARK)
]

for idx, (title, body, accent) in enumerate(recs):
    top = Inches(1.5 + idx * 1.35)
    add_card(s9, Inches(0.8), top, Inches(11.7), Inches(1.2), bg_color=RGBColor(255, 255, 255), border_color=accent)
    tb = s9.shapes.add_textbox(Inches(1.1), top + Inches(0.1), Inches(11.1), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = f"{idx + 1}. {title}"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = accent
    
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(10)
    p2.font.color.rgb = C_TEXT_DARK
    p2.space_before = Pt(4)

# ==========================================
# SLIDE 10: INTERVIEW QUESTIONS SUMMARY
# ==========================================
s10 = prs.slides.add_slide(blank_layout)
add_header(s10, "Technical Viva & Interview Questions Reference", "INTERVIEW PREPARATION")

qa_pairs = [
    ("Q1: Key Elements of a Dashboard?", 
     "KPI Summary Cards (totals), Trend/Time-series visuals, Slicers/Filters for interactivity, Category/Distribution charts, consistent color hierarchy, and intuitive navigation."),
    ("Q2: What is a KPI?", 
     "A Key Performance Indicator is a quantifiable metric reflecting how effectively an organization is achieving its business goals (e.g., Sales, Net Margin, YoY Growth)."),
    ("Q3: Slicers in Power BI / Tableau?", 
     "Interactive visual filters placed directly on the canvas allowing end-users to slice data dynamically across dimensions (Region, Date, Category)."),
    ("Q4: Difference between Power BI & Tableau?", 
     "Tableau excels in deep statistical exploration and custom visual flexibility; Power BI offers seamless Microsoft 365 integration, tabular DAX modeling, and lower cost."),
    ("Q5: How to make a dashboard interactive?", 
     "Through Filter Slicers, Parameter Actions, Cross-Highlighting, URL navigation actions, Dynamic tooltips, and drill-through pages."),
    ("Q6: Dealing with large datasets?", 
     "Use Tableau Data Extracts (.hyper), pre-aggregation, pushdown SQL queries, indexing, incremental refreshes, and hide unused columns."),
    ("Q7: Best chart types for trend analysis?", 
     "Line charts (standard for time-series), Area charts (cumulative volume), Dual-axis lines (Sales vs Profit), and Sparklines (micro trends).")
]

for idx, (q, a) in enumerate(qa_pairs):
    row = idx // 2
    col = idx % 2
    left = Inches(0.8 + col * 5.95)
    top = Inches(1.5 + row * 1.35)
    
    add_card(s10, left, top, Inches(5.75), Inches(1.22), bg_color=RGBColor(248, 250, 252), border_color=RGBColor(203, 213, 225))
    tb = s10.shapes.add_textbox(left + Inches(0.15), top + Inches(0.08), Inches(5.45), Inches(1.05))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p_q = tf.paragraphs[0]
    p_q.text = q
    p_q.font.size = Pt(10.5)
    p_q.font.bold = True
    p_q.font.color.rgb = C_NAVY_BLUE
    
    p_a = tf.add_paragraph()
    p_a.text = a
    p_a.font.size = Pt(8.5)
    p_a.font.color.rgb = C_TEXT_DARK
    p_a.space_before = Pt(3)

pptx_path = 'deliverables/Executive_Sales_Dashboard_Summary.pptx'
prs.save(pptx_path)
print(f"Presentation successfully created and saved at: {pptx_path}")
